from io import BytesIO
import pandas as pd
from openpyxl import Workbook,load_workbook
from src.workpaper_filler import append_rationales_column,fill_workpaper

def make_book(value=None):
    wb=Workbook();ws=wb.active;ws.title="可比公司法";ws["E11"]=value
    buf=BytesIO();wb.save(buf);return buf.getvalue()

def reviewed(cell="E11",company="虚构甲公司",decision="最终纳入"):
    return pd.DataFrame([{"investee_company_name":"虚构被投企业","company_name":company,"人工最终决定":decision,"合理性分析初稿":"相似点与差异点均来自虚构输入资料。","workpaper_file":"审计底稿.xlsx","sheet_name":"可比公司法","cell_address":cell}])

def test_writes_only_to_output_copy():
    original=make_book();filled,report=fill_workpaper(original,"审计底稿.xlsx",reviewed())
    assert report.iloc[0]["状态"]=="已写入"
    assert load_workbook(BytesIO(filled))["可比公司法"]["E11"].value.startswith("相似点")
    assert load_workbook(BytesIO(original))["可比公司法"]["E11"].value is None

def test_requires_confirmation_to_overwrite_nonempty_cell():
    _,skipped=fill_workpaper(make_book("原有底稿内容"),"审计底稿.xlsx",reviewed())
    filled,allowed=fill_workpaper(make_book("原有底稿内容"),"审计底稿.xlsx",reviewed(),overwrite=True)
    assert skipped.iloc[0]["状态"]=="跳过"
    assert allowed.iloc[0]["状态"]=="已写入"
    assert load_workbook(BytesIO(filled))["可比公司法"]["E11"].value.startswith("相似点")

def test_duplicate_cell_mappings_are_rejected():
    df=pd.concat([reviewed(),reviewed(company="虚构乙公司")],ignore_index=True)
    _,report=fill_workpaper(make_book(),"审计底稿.xlsx",df)
    assert (report["状态"]=="跳过").all()
    assert report["说明"].str.contains("重复").all()

def test_non_included_rows_are_explained_as_not_written():
    df=reviewed(decision="待进一步讨论")
    _,report=fill_workpaper(make_book(),"审计底稿.xlsx",df)
    assert report.iloc[0]["状态"]=="未写入"
    assert "仅最终纳入" in report.iloc[0]["说明"]

def test_exports_only_reviewed_and_included_pairs_to_clean_final_sheet():
    wb=Workbook(); ws=wb.active; ws.title="client_list"
    ws.append(["investee_company_name","comparable_company_name"])
    ws.append(["被投甲","客户可比公司"])
    buf=BytesIO(); wb.save(buf); original=buf.getvalue()
    results=pd.DataFrame([
        {"investee_company_name":"被投甲","company_name":"程序候选一","stock_code":"600001","人工复核状态":"已复核","人工最终结论":"纳入","合理性分析初稿":"两家公司业务模式及客户群体相近，因此具备可比性。"},
        {"investee_company_name":"被投甲","company_name":"程序候选二","stock_code":"600002","人工复核状态":"待复核","人工最终结论":"纳入","合理性分析初稿":"待复核记录不应写入。"},
        {"investee_company_name":"被投甲","company_name":"未纳入公司","stock_code":"600003","人工复核状态":"已复核","人工最终结论":"不纳入","合理性分析初稿":"排除记录不应写入。"},
    ])
    filled,report=append_rationales_column(original,"input.xlsx",results)
    book=load_workbook(BytesIO(filled))
    out=book["最终可比公司名单"]
    headers=[out.cell(1,c).value for c in range(1,out.max_column+1)]
    assert headers==["被投企业","最终可比公司","证券代码","最终合理性分析"]
    rows=list(out.iter_rows(min_row=2,values_only=True))
    assert rows==[("被投甲","程序候选一","600001","两家公司业务模式及客户群体相近，因此具备可比性。")]
    assert set(report["状态"])=={"已纳入"}
    assert book["client_list"].max_column==2
    assert load_workbook(BytesIO(original))["client_list"].max_column==2

def test_replaces_existing_final_sheet_without_touching_source_table():
    wb=Workbook(); ws=wb.active; ws.title="客户名单"
    ws.append(["被投企业","可比公司"]); ws.append(["被投甲","候选甲"])
    final=wb.create_sheet("最终可比公司名单"); final.append(["旧字段"]); final.append(["旧内容"])
    buf=BytesIO(); wb.save(buf)
    results=pd.DataFrame([{"investee_company_name":"被投甲","company_name":"候选甲","stock_code":"000001","人工复核状态":"已复核","人工最终结论":"纳入","合理性分析初稿":"已有客户候选经分析后可纳入。"}])
    filled,report=append_rationales_column(buf.getvalue(),"input.xlsx",results)
    book=load_workbook(BytesIO(filled)); out=book["最终可比公司名单"]
    assert out.max_row==2
    assert out.cell(2,2).value=="候选甲"
    assert out.cell(2,4).value=="已有客户候选经分析后可纳入。"
    assert book["客户名单"].max_column==2
    assert report.iloc[0]["状态"]=="已纳入"
    assert book.active.title=="最终可比公司名单"

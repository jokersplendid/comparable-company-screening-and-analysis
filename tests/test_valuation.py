from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.data_loader import create_input_template, read_workbook, build_tables_from_mapping
from src.valuation import (enterprise_value, peer_multiples, build_peer_valuation_detail,
                           multiple_statistics, target_valuation, valuation_summary,
                           valuation_calculation_detail, apply_valuation_plan)
from src.workpaper_filler import append_valuation_sheets


def test_enterprise_value_uses_interest_bearing_debt_not_total_liabilities():
    record={"market_cap":500,"cash":40,"interest_bearing_debt":100,"total_liabilities":900,
            "minority_interest":10,"preferred_stock":5,"non_operating_investments":15}
    assert enterprise_value(record)==560


def test_peer_multiples_and_target_equity_bridge():
    record={"market_cap":500,"cash":40,"interest_bearing_debt":100,"revenue":100,"ebitda":50,
            "net_profit":25,"book_equity":200}
    multiples=peer_multiples(record)
    assert multiples["企业价值"]==560
    assert multiples["EV/Revenue"]==5.6
    assert multiples["EV/EBITDA"]==11.2
    stats=pd.DataFrame([{"被投企业":"目标","估值方法":"EV/EBITDA","有效样本数":3,
                         "P25":8,"中位数":10,"P75":12}])
    result=target_valuation({"company_name":"目标","ebitda":20,"cash":10,"interest_bearing_debt":30},stats)
    assert result.iloc[0]["中值股权价值"]==180


def test_valuation_calculation_detail_explains_equity_bridge_and_adjustments():
    stats=pd.DataFrame([{"被投企业":"目标","估值方法":"EV/EBITDA","有效样本数":3,
                         "P25":8,"中位数":10,"P75":12}])
    result=target_valuation(
        {"company_name":"目标","ebitda":20,"cash":10,"interest_bearing_debt":30},
        stats,control_premium=.10,illiquidity_discount=.20)
    detail=valuation_calculation_detail(result,control_premium=.10,illiquidity_discount=.20)
    middle=detail.loc[detail["情景"].eq("中值（中位数）")].iloc[0]
    assert middle["采用倍数"]==10
    assert middle["隐含企业价值"]==200
    assert middle["企业价值转股权调整"]==-20
    assert middle["调整前股权价值"]==180
    assert middle["最终股权价值"]==pytest.approx(158.4)


def test_custom_valuation_plan_uses_selected_method_and_multiple():
    plan=pd.DataFrame([{"估值方法":"EV/EBITDA","情景":"自定义倍数","采用倍数":9.5}])
    detail=apply_valuation_plan(
        {"company_name":"目标","ebitda":20,"cash":10,"interest_bearing_debt":30},plan)
    row=detail.iloc[0]
    assert row["估值方法"]=="EV/EBITDA"
    assert row["隐含企业价值"]==190
    assert row["最终股权价值"]==170


def test_custom_valuation_method_supports_custom_value_basis_and_metric():
    plan=pd.DataFrame([{"估值方法":"EV/总资产","情景":"自定义倍数","采用倍数":1.8}])
    detail=apply_valuation_plan(
        {"company_name":"目标","total_assets":200,"cash":20,"interest_bearing_debt":50},plan,
        custom_methods={"EV/总资产":{"metric":"total_assets","basis":"EV","formula":"企业价值 / 总资产"}})
    row=detail.iloc[0]
    assert row["估值公式"]=="企业价值 / 总资产"
    assert row["隐含企业价值"]==360
    assert row["最终股权价值"]==330


def test_custom_adjustments_are_applied_at_ev_and_equity_stages():
    plan=pd.DataFrame([{"估值方法":"EV/EBITDA","情景":"P50","采用倍数":10}])
    adjustments=pd.DataFrame([
        {"调整项名称":"非经营性调整","调整阶段":"企业价值","调整类型":"固定金额","运算方式":"加","数值":20},
        {"调整项名称":"流动性折价","调整阶段":"股权价值","调整类型":"百分比","运算方式":"减","数值":10},
    ])
    detail=apply_valuation_plan(
        {"company_name":"目标","ebitda":20,"cash":10,"interest_bearing_debt":30},plan,
        adjustments=adjustments)
    row=detail.iloc[0]
    assert row["隐含企业价值"]==220
    assert row["调整前股权价值"]==200
    assert row["自定义调整后股权价值"]==180
    assert row["最终股权价值"]==180
    assert "非经营性调整" in row["自定义调整"]
    assert "流动性折价" in row["自定义调整"]


def test_valuation_pipeline_groups_peers_by_investee():
    included=pd.DataFrame([
        {"investee_company_name":"目标甲","company_name":"可比一"},
        {"investee_company_name":"目标甲","company_name":"可比二"},
    ])
    candidates=pd.DataFrame([
        {"company_name":"可比一","market_cap":100,"cash":10,"interest_bearing_debt":20,"revenue":50,"ebitda":10,"net_profit":5,"book_equity":40},
        {"company_name":"可比二","market_cap":150,"cash":15,"interest_bearing_debt":30,"revenue":60,"ebitda":12,"net_profit":6,"book_equity":50},
    ])
    detail=build_peer_valuation_detail(included,candidates)
    stats=multiple_statistics(detail)
    result=target_valuation({"company_name":"目标甲","revenue":55,"ebitda":11,"net_profit":5.5,"book_equity":45,"cash":12,"interest_bearing_debt":25},stats)
    summary=valuation_summary(result)
    assert set(result["估值方法"])=={"EV/Revenue","EV/EBITDA","P/E","P/B"}
    assert summary.iloc[0]["有效估值方法数"]==4


def test_template_can_be_loaded_without_column_mapping():
    template=create_input_template()
    tables=read_workbook(BytesIO(template))
    assert set(tables)=={"investee","candidates","client_list"}
    assert "interest_bearing_debt" in tables["investee"].columns
    assert "market_cap" in tables["candidates"].columns


def test_custom_mapping_requires_every_required_field():
    raw={"目标":pd.DataFrame({"公司":["A"]}),"候选":pd.DataFrame({"公司":["B"]}),"名单":pd.DataFrame({"目标":["A"],"可比":["B"]})}
    try:
        build_tables_from_mapping(raw,{"investee":"目标","candidates":"候选","client_list":"名单"},{"investee":{"company_name":"公司"},"candidates":{"company_name":"公司"},"client_list":{"investee_company_name":"目标","comparable_company_name":"可比"}})
    except ValueError as exc:
        assert "尚未映射必需字段" in str(exc)
    else:
        raise AssertionError("missing mappings should be rejected")


def test_complete_export_adds_valuation_sheets_to_original_copy():
    source=create_input_template()
    reviewed=pd.DataFrame([{"investee_company_name":"目标甲","company_name":"可比一","stock_code":"600001",
                            "人工复核状态":"已复核","人工最终结论":"纳入","合理性分析初稿":"业务和盈利模式相近。"}])
    sheets={"可比公司估值明细":pd.DataFrame([{"被投企业":"目标甲","可比公司":"可比一","EV/EBITDA":10.0}]),
            "估值汇总":pd.DataFrame([{"被投企业":"目标甲","股权价值中值":100.0}])}
    output,_=append_valuation_sheets(source,"input.xlsx",reviewed,sheets)
    wb=load_workbook(BytesIO(output))
    assert "最终可比公司名单" in wb.sheetnames
    assert "可比公司估值明细" in wb.sheetnames
    assert wb["估值汇总"]["B2"].value==100

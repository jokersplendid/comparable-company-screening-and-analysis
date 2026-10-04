from io import BytesIO
import json
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def export_results(sheets: dict[str,pd.DataFrame], run_log: list[dict]) -> bytes:
    output=BytesIO()
    with pd.ExcelWriter(output,engine="openpyxl") as writer:
        for name,df in sheets.items(): df.to_excel(writer,sheet_name=name[:31],index=False)
        pd.DataFrame(run_log).to_excel(writer,sheet_name="运行日志",index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="183B56");cell.alignment=Alignment(wrap_text=True)
            for col in ws.columns:
                width=min(48,max(12,max((len(str(c.value or "")) for c in list(col)[:100]),default=12)+2))
                ws.column_dimensions[get_column_letter(col[0].column)].width=width
    return output.getvalue()

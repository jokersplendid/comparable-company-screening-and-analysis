from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

REQUIRED = {
 "investee": ["company_id", "company_name", "industry", "company_description", "main_business", "products", "application_scenarios", "customer_groups", "business_model", "revenue", "total_assets", "revenue_growth", "gross_margin", "ebitda_margin"],
 "candidates": ["company_name", "industry", "company_description", "main_business", "products", "application_scenarios", "customer_groups", "business_model", "revenue", "total_assets", "revenue_growth", "gross_margin", "ebitda_margin"],
 "client_list": ["investee_company_name", "comparable_company_name"],
}
OPTIONAL = ["stock_code", "market", "market_cap", "valuation_multiple", "valuation_multiple_type", "source_reference", "total_liabilities", "debt_to_assets", "cash", "interest_bearing_debt", "minority_interest", "preferred_stock", "non_operating_investments", "ebitda", "net_profit", "book_equity", "reporting_period", "valuation_date", "currency", "amount_unit", "workpaper_file", "sheet_name", "cell_address", "industry_subcode", "industry_code", "csrc_industry_code", "wind_industry_code", "industry_major_code", "industry_category_code", "industry_major", "industry_category", "related_industries", "adjacent_industries", "upstream_industries", "downstream_industries"]

FIELD_LABELS = {
 "company_id":"公司ID", "company_name":"公司名称", "industry":"行业", "company_description":"公司简介",
 "main_business":"主营业务", "products":"产品", "application_scenarios":"应用场景", "customer_groups":"客户群体",
 "business_model":"商业模式", "revenue":"营业收入", "total_assets":"总资产", "total_liabilities":"总负债",
 "revenue_growth":"收入增长率", "gross_margin":"毛利率", "ebitda_margin":"EBITDA利润率", "stock_code":"证券代码",
 "market":"市场", "market_cap":"市值", "cash":"现金及现金等价物", "interest_bearing_debt":"有息负债",
 "minority_interest":"少数股东权益", "preferred_stock":"优先股", "non_operating_investments":"非经营性投资",
 "ebitda":"EBITDA", "net_profit":"净利润", "book_equity":"账面净资产", "reporting_period":"报告期",
 "valuation_date":"估值基准日", "currency":"币种", "amount_unit":"金额单位",
 "investee_company_name":"被投企业", "comparable_company_name":"客户提供的可比公司",
}

COLUMN_ALIASES = {
 "公司id":"company_id", "公司编号":"company_id", "企业id":"company_id", "公司名称":"company_name", "企业名称":"company_name",
 "被投企业":"investee_company_name", "被投企业名称":"investee_company_name", "可比公司":"comparable_company_name",
 "候选公司":"comparable_company_name", "行业":"industry", "所属行业":"industry", "公司简介":"company_description",
 "主营业务":"main_business", "产品":"products", "主要产品":"products", "应用场景":"application_scenarios",
 "客户群体":"customer_groups", "商业模式":"business_model", "营业收入":"revenue", "营收":"revenue",
 "总资产":"total_assets", "总负债":"total_liabilities", "收入增长率":"revenue_growth", "营收增长率":"revenue_growth",
 "毛利率":"gross_margin", "ebitda利润率":"ebitda_margin", "ebitda率":"ebitda_margin", "证券代码":"stock_code",
 "股票代码":"stock_code", "市场":"market", "上市市场":"market", "市值":"market_cap", "总市值":"market_cap",
 "现金":"cash", "现金及现金等价物":"cash", "有息负债":"interest_bearing_debt", "带息负债":"interest_bearing_debt",
 "少数股东权益":"minority_interest", "优先股":"preferred_stock", "非经营性投资":"non_operating_investments",
 "ebitda":"ebitda", "净利润":"net_profit", "归母净利润":"net_profit", "账面净资产":"book_equity",
 "归母净资产":"book_equity", "报告期":"reporting_period", "估值基准日":"valuation_date", "币种":"currency", "金额单位":"amount_unit",
}

def _norm(value) -> str:
    return str(value or "").strip().casefold().replace(" ", "").replace("_", "")

def read_raw_workbook(source) -> dict[str, pd.DataFrame]:
    xls = pd.ExcelFile(source)
    return {sheet: pd.read_excel(xls, sheet_name=sheet) for sheet in xls.sheet_names}

def suggest_column_mapping(frame: pd.DataFrame, table: str) -> dict[str, str]:
    reverse={_norm(field):field for field in REQUIRED[table]+OPTIONAL}
    reverse.update({_norm(alias):field for alias,field in COLUMN_ALIASES.items()})
    mapping={}
    for column in frame.columns:
        field=reverse.get(_norm(column))
        if field and field in REQUIRED[table]+OPTIONAL and field not in mapping:
            mapping[field]=str(column)
    return mapping

def build_tables_from_mapping(raw: dict[str,pd.DataFrame], sheet_mapping: dict[str,str],
                              column_mapping: dict[str,dict[str,str]]) -> dict[str,pd.DataFrame]:
    result={}
    for table in REQUIRED:
        sheet=sheet_mapping.get(table)
        if not sheet or sheet not in raw:
            raise ValueError(f"请为 {table} 选择工作表。")
        source=raw[sheet]
        mapping=column_mapping.get(table,{})
        missing=[field for field in REQUIRED[table] if not mapping.get(field)]
        if missing:
            labels="、".join(FIELD_LABELS.get(field,field) for field in missing)
            raise ValueError(f"{table} 尚未映射必需字段：{labels}")
        renamed={source_col:field for field,source_col in mapping.items() if source_col in source.columns}
        result[table]=source.rename(columns=renamed).copy()
    return result

def create_input_template() -> bytes:
    wb=Workbook(); wb.remove(wb.active)
    for table in ["investee","candidates","client_list"]:
        ws=wb.create_sheet(table)
        fields=list(REQUIRED[table])
        if table in {"investee","candidates"}:
            fields += [field for field in OPTIONAL if field not in fields]
        for col,field in enumerate(fields,1):
            cell=ws.cell(1,col,field)
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="183B56")
            cell.alignment=Alignment(wrap_text=True)
            cell.comment=Comment(FIELD_LABELS.get(field,field),"系统字段说明")
            ws.column_dimensions[cell.column_letter].width=max(14,min(24,len(field)+3))
        ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{ws.cell(1,len(fields)).column_letter}1"
    output=BytesIO(); wb.save(output); return output.getvalue()

def read_workbook(source) -> dict[str, pd.DataFrame]:
    xls = pd.ExcelFile(source)
    aliases = {"被投企业": "investee", "Investee": "investee", "investee": "investee", "候选公司": "candidates", "上市公司候选池": "candidates", "Candidates": "candidates", "candidates": "candidates", "客户名单": "client_list", "客户已选名单": "client_list", "Client_List": "client_list", "client_list": "client_list"}
    found = {}
    for sheet in xls.sheet_names:
        key = aliases.get(sheet)
        if key:
            frame=pd.read_excel(xls, sheet_name=sheet)
            mapping=suggest_column_mapping(frame,key)
            frame=frame.rename(columns={source_col:field for field,source_col in mapping.items()})
            found[key] = frame
    missing = set(REQUIRED) - set(found)
    if missing: raise ValueError(f"缺少工作表：{', '.join(sorted(missing))}。请使用 investee、candidates、client_list 三个工作表。")
    return found

def quality_report(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows=[]
    for table, df in tables.items():
        required=REQUIRED[table]
        for col in required:
            if col not in df.columns:
                rows.append({"数据表":table,"字段":col,"问题": "缺少必需列","行数":len(df)})
            else:
                n=int(df[col].isna().sum() + (df[col].astype(str).str.strip()=="").sum())
                if n: rows.append({"数据表":table,"字段":col,"问题":"字段为空","行数":n})
        for col in OPTIONAL:
            if col in df.columns:
                n=int(df[col].isna().sum())
                if n: rows.append({"数据表":table,"字段":col,"问题":"可选字段为空","行数":n})
    return pd.DataFrame(rows, columns=["数据表","字段","问题","行数"])
